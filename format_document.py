from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


SOURCE = Path('/home/neo/Downloads/Marks_Analyser_Project_Documentation (1).docx')
OUTPUT = Path('/home/neo/Desktop/marks-analyser/Marks_Analyser_Project_Documentation_Revised.docx')
CODE_TABLES = range(19, 26)
SCREENSHOT_TABLES = set(range(14, 18)) | set(range(26, 35))
MAX_SCREENSHOT_WIDTH = Inches(6.05)


def set_cell_borders_none(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        tc_pr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        tag = qn(f'w:{edge}')
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f'w:{edge}')
            borders.append(element)
        element.set(qn('w:val'), 'nil')


def clear_table_borders(table):
    for row in table.rows:
        for cell in row.cells:
            set_cell_borders_none(cell)


def resize_and_center_images(table):
    clear_table_borders(table)
    for paragraph in table._tbl.xpath('.//w:p'):
        blips = paragraph.xpath('.//a:blip')
        if not blips:
            continue
        paragraph.set(qn('w:jc'), 'center')
        for extent in paragraph.xpath('.//wp:extent'):
            current_width = int(extent.get('cx'))
            current_height = int(extent.get('cy'))
            if current_width > MAX_SCREENSHOT_WIDTH:
                scale = MAX_SCREENSHOT_WIDTH / current_width
                extent.set('cx', str(int(MAX_SCREENSHOT_WIDTH)))
                extent.set('cy', str(int(current_height * scale)))
                # Keep the drawing's visible bounds in sync with its layout extent.
                for graphic_extent in paragraph.xpath('.//a:xfrm/a:ext'):
                    graphic_extent.set('cx', str(int(MAX_SCREENSHOT_WIDTH)))
                    graphic_extent.set('cy', str(int(current_height * scale)))


def resize_and_center_body_images(doc):
    for paragraph in doc.paragraphs:
        extents = paragraph._p.xpath('.//wp:extent')
        if not extents:
            continue
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for extent in extents:
            current_width = int(extent.get('cx'))
            current_height = int(extent.get('cy'))
            if current_width > MAX_SCREENSHOT_WIDTH:
                scale = MAX_SCREENSHOT_WIDTH / current_width
                extent.set('cx', str(int(MAX_SCREENSHOT_WIDTH)))
                extent.set('cy', str(int(current_height * scale)))
                for graphic_extent in paragraph._p.xpath('.//a:xfrm/a:ext'):
                    graphic_extent.set('cx', str(int(MAX_SCREENSHOT_WIDTH)))
                    graphic_extent.set('cy', str(int(current_height * scale)))


def append_page_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement('w:fldChar')
    begin.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = ' PAGE '
    separate = OxmlElement('w:fldChar')
    separate.set(qn('w:fldCharType'), 'separate')
    text = OxmlElement('w:t')
    text.text = '1'
    end = OxmlElement('w:fldChar')
    end.set(qn('w:fldCharType'), 'end')
    run._r.extend([begin, instr, separate, text, end])


def add_page_border_and_number(section):
    sect_pr = section._sectPr
    old = sect_pr.find(qn('w:pgBorders'))
    if old is not None:
        sect_pr.remove(old)
    borders = OxmlElement('w:pgBorders')
    borders.set(qn('w:offsetFrom'), 'page')
    for edge in ('top', 'left', 'bottom', 'right'):
        border = OxmlElement(f'w:{edge}')
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), '8')
        border.set(qn('w:space'), '24')
        border.set(qn('w:color'), '000000')
        borders.append(border)
    sect_pr.append(borders)

    section.different_first_page_header_footer = False
    header = section.header
    paragraph = header.paragraphs[0]
    paragraph.clear()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.paragraph_format.space_after = Pt(0)
    append_page_field(paragraph)


def insert_code_as_text(table):
    lines = []
    for paragraph in table._tbl.xpath('.//w:p'):
        text = ''.join(paragraph.itertext())
        lines.extend(text.splitlines() or [''])

    anchor = table._tbl
    for line in lines:
        new_p = OxmlElement('w:p')
        p_pr = OxmlElement('w:pPr')
        spacing = OxmlElement('w:spacing')
        spacing.set(qn('w:before'), '0')
        spacing.set(qn('w:after'), '0')
        spacing.set(qn('w:line'), '200')
        spacing.set(qn('w:lineRule'), 'auto')
        p_pr.append(spacing)
        new_p.append(p_pr)
        run = OxmlElement('w:r')
        r_pr = OxmlElement('w:rPr')
        font = OxmlElement('w:rFonts')
        font.set(qn('w:ascii'), 'Consolas')
        font.set(qn('w:hAnsi'), 'Consolas')
        r_pr.append(font)
        size = OxmlElement('w:sz')
        size.set(qn('w:val'), '17')
        r_pr.append(size)
        run.append(r_pr)
        txt = OxmlElement('w:t')
        if line.startswith(' ') or line.endswith(' '):
            txt.set(qn('xml:space'), 'preserve')
        txt.text = line
        run.append(txt)
        new_p.append(run)
        anchor.addnext(new_p)
        anchor = new_p
    table._tbl.getparent().remove(table._tbl)


def main():
    doc = Document(SOURCE)
    tables = list(doc.tables)
    for index, table in enumerate(tables):
        if index in SCREENSHOT_TABLES:
            resize_and_center_images(table)
    resize_and_center_body_images(doc)
    for index in reversed(list(CODE_TABLES)):
        insert_code_as_text(tables[index])
    for section in doc.sections:
        add_page_border_and_number(section)
    update_fields = doc.settings.element.find(qn('w:updateFields'))
    if update_fields is None:
        update_fields = OxmlElement('w:updateFields')
        doc.settings.element.append(update_fields)
    update_fields.set(qn('w:val'), 'true')
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == '__main__':
    main()
